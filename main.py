import os, csv, io, json, time, sqlite3, secrets, hmac, hashlib, re, base64
from contextlib import closing
from difflib import SequenceMatcher
from collections import defaultdict
from pathlib import Path

import httpx
from openpyxl import load_workbook, Workbook
from fastapi import FastAPI, UploadFile, File, HTTPException, Request, Form
from fastapi.responses import HTMLResponse, Response
from fastapi.templating import Jinja2Templates

DEV_MODE        = os.environ.get("DEV_MODE", "1") == "1"
PAYSTACK_SECRET = os.environ.get("PAYSTACK_SECRET_KEY", "")
BREVO_API_KEY   = os.environ.get("BREVO_API_KEY", "")
FROM_EMAIL      = os.environ.get("FROM_EMAIL", "bsquaressoftwares@gmail.com")
FROM_NAME       = os.environ.get("FROM_NAME", "VendorCheck")
PRICE_KOBO      = int(os.environ.get("PRICE_KOBO", "50000"))
PRICE_DISPLAY   = os.environ.get("PRICE_DISPLAY", "500")
DB_PATH         = os.environ.get("DB_PATH", "data/vendorcheck.db")
PAY_EMAIL       = os.environ.get("PAY_EMAIL", "bsquaressoftwares@gmail.com")

Path(DB_PATH).parent.mkdir(parents=True, exist_ok=True)

app = FastAPI()
templates = Jinja2Templates(directory="templates")

def init_db():
    with closing(sqlite3.connect(DB_PATH)) as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS orders (
            checkout_id TEXT PRIMARY KEY, filename TEXT, total INTEGER, flagged INTEGER,
            fraud INTEGER, cleaned_csv TEXT, paid INTEGER DEFAULT 0, email TEXT,
            phone TEXT, reference TEXT, created REAL)""")
        conn.commit()
init_db()

def db():
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    return c

def purge_old():
    cutoff = time.time() - 86400
    with closing(db()) as conn:
        conn.execute("DELETE FROM orders WHERE created < ?", (cutoff,))
        conn.commit()

def read_csv_records(raw_bytes):
    text = raw_bytes.decode("utf-8", errors="ignore")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(r) for r in reader]

def read_xlsx_records(raw_bytes):
    wb = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(it)]
    except StopIteration:
        return []
    out = []
    for row in it:
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        d = {}
        for i, h in enumerate(headers):
            if i < len(row):
                v = row[i]
                d[h] = "" if v is None else str(v)
        out.append(d)
    return out

def records_from_upload(filename, raw_bytes):
    fn = (filename or "").lower()
    if fn.endswith(".xlsx") or fn.endswith(".xls"):
        return read_xlsx_records(raw_bytes)
    if fn.endswith(".csv"):
        return read_csv_records(raw_bytes)
    try:
        return read_csv_records(raw_bytes)
    except Exception:
        return read_xlsx_records(raw_bytes)

TIN_RE = re.compile(r"^[A-Z]{1,2}\d{6,9}[A-Z]$")
COLS = {
    "name": ["name","vendor","vendor_name","supplier","supplier_name","company"],
    "tin":  ["tin","kra_pin","pin","kra","tax_id"],
    "bank": ["bank_account","bank_acc","account","account_no","acc","bank"],
    "phone":["phone","telephone","msisdn","mobile","tel"],
    "email":["email","e_mail","mail"],
}
def pick(row, keys):
    for k in row:
        if k and k.strip().lower() in keys: return row[k] or ""
    return ""
def nname(s): return re.sub(r"[^a-z0-9 ]", "", (s or "").lower()).strip()
def ntin(s):  return re.sub(r"[^A-Z0-9]", "", (s or "").upper())
def ndig(s):  return re.sub(r"\D", "", s or "")
def phone_ok(d):
    if not d: return False
    return ((len(d)==12 and d.startswith("254")) or
            (len(d)==10 and d.startswith("0")) or
            (len(d)==9  and d[0] in ("7","1")))
def ratio(a, b):
    if a == b: return 1.0
    if not a or not b: return 0.0
    return SequenceMatcher(None, a, b).ratio()

def clean_records(records):
    rows = []
    for r in records:
        rows.append({
            "name":  pick(r, COLS["name"]).strip(),
            "tin":   ntin(pick(r, COLS["tin"])),
            "bank":  ndig(pick(r, COLS["bank"])),
            "phone": ndig(pick(r, COLS["phone"])),
            "email": pick(r, COLS["email"]).strip(),
        })
    if not rows:
        raise ValueError("No data rows found in the file.")

    tin_map = defaultdict(list); bank_map = defaultdict(list)
    for i, r in enumerate(rows):
        if r["tin"]:  tin_map[r["tin"]].append(i)
        if r["bank"]: bank_map[r["bank"]].append(i)

    order = sorted(range(len(rows)), key=lambda i: nname(rows[i]["name"]))
    near = set()
    for a, b in zip(order, order[1:]):
        if rows[a]["name"] and rows[b]["name"] and \
           ratio(nname(rows[a]["name"]), nname(rows[b]["name"])) > 0.90:
            near.add(a); near.add(b)

    for i, r in enumerate(rows):
        f = []
        if not r["tin"]: f.append("MISSING_TIN")
        elif not TIN_RE.match(r["tin"]): f.append("INVALID_TIN_FORMAT")
        elif len(tin_map[r["tin"]]) > 1: f.append("TIN_SHARED")

        if not r["bank"]: f.append("MISSING_BANK")
        else:
            if not 6 <= len(r["bank"]) <= 17: f.append("ODD_BANK_LENGTH")
            same = bank_map[r["bank"]]
            if len(same) > 1:
                others = [rows[j] for j in same if j != i]
                if any(o["tin"] and o["tin"] != r["tin"] for o in others):
                    f.append("SAME_BANK_DIFF_TIN")
                else:
                    f.append("DUPLICATE_BANK")

        if not r["phone"]: f.append("MISSING_PHONE")
        elif not phone_ok(r["phone"]): f.append("ODD_PHONE")
        if i in near: f.append("NEAR_DUPLICATE_NAME")
        r["flags"] = ";".join(f) if f else "OK"

    flagged = [r for r in rows if r["flags"] != "OK"]
    fraud   = [r for r in rows if "SAME_BANK_DIFF_TIN" in r["flags"]]
    summary = {
        "total": len(rows), "flagged": len(flagged), "clean": len(rows)-len(flagged),
        "fraud": len(fraud),
        "bad_tin": sum(1 for r in rows if "INVALID_TIN_FORMAT" in r["flags"] or "MISSING_TIN" in r["flags"]),
        "dup_bank": sum(1 for r in rows if "DUPLICATE_BANK" in r["flags"]),
        "near": sum(1 for r in rows if "NEAR_DUPLICATE_NAME" in r["flags"]),
    }
    return rows, summary, flagged, fraud

def rows_to_csv(rows):
    out = io.StringIO()
    w = csv.DictWriter(out, fieldnames=["name","tin","bank","phone","email","flags"])
    w.writeheader()
    for r in rows: w.writerow({k: r[k] for k in w.fieldnames})
    return out.getvalue()

async def send_csv_email(to_email, original_filename, csv_content, stats):
    if not BREVO_API_KEY:
        print("BREVO: no key set, skipping email")
        return False
    b64 = base64.b64encode(csv_content.encode("utf-8")).decode("ascii")
    stem = original_filename.rsplit(".", 1)[0] if "." in original_filename else original_filename
    clean_name = "cleaned_" + stem + ".csv"
    html = (
        "<p>Your cleaned vendor file is attached.</p>"
        "<p><strong>" + str(stats.get("total", "")) + "</strong> vendors scanned - "
        "<strong>" + str(stats.get("flagged", "")) + "</strong> flagged - "
        "<strong>" + str(stats.get("fraud", "")) + "</strong> fraud signals</p>"
        "<p>Keep this email for your audit trail. The file contains a <code>flags</code> column "
        "you can filter in Excel or Google Sheets.</p>"
        "<p>- VendorCheck</p>"
    )
    payload = {
        "sender": {"email": FROM_EMAIL, "name": FROM_NAME},
        "to": [{"email": to_email}],
        "subject": "Your cleaned vendor file - " + clean_name,
        "htmlContent": html,
        "attachment": [{"content": b64, "name": clean_name}],
    }
    try:
        async with httpx.AsyncClient(timeout=30) as client:
            r = await client.post(
                "https://api.brevo.com/v3/smtp/email",
                headers={"api-key": BREVO_API_KEY, "Content-Type": "application/json"},
                json=payload,
            )
            print("BREVO status:", r.status_code, "body:", r.text[:500])
            return r.status_code in (200, 201, 202)
    except Exception as e:
        print("BREVO exception:", repr(e))
        return False

@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    return templates.TemplateResponse("home.html", {
        "request": request, "dev_mode": DEV_MODE, "pay_email": PAY_EMAIL,
    })

@app.get("/template.xlsx")
def template_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendors"
    ws.append(["name", "tin", "bank_account", "phone", "email"])
    ws.append(["Acme Supplies Ltd", "P051234567X", "0123456789", "0712345678", "info@acme.co.ke"])
    ws.append(["Beta Traders", "P051234567Y", "0987654321", "0700000000", "b@beta.co.ke"])
    ws.append(["Gamma Ltd", "P051234568Z", "1122334455", "0711111111", "g@gamma.co.ke"])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="vendor_template.xlsx"'},
    )

@app.post("/scan", response_class=HTMLResponse)
async def scan(request: Request, file: UploadFile = File(...)):
    purge_old()
    raw = await file.read()
    try:
        records = records_from_upload(file.filename, raw)
        rows, summary, flagged, fraud = clean_records(records)
    except Exception as e:
        return templates.TemplateResponse("error.html", {
            "request": request, "msg": str(e), "dev_mode": DEV_MODE,
        }, status_code=400)

    checkout_id = secrets.token_urlsafe(12)
    with closing(db()) as conn:
        conn.execute("INSERT INTO orders (checkout_id, filename, total, flagged, fraud, cleaned_csv, created) VALUES (?,?,?,?,?,?,?)",
            (checkout_id, file.filename, summary["total"], summary["flagged"],
             summary["fraud"], rows_to_csv(rows), time.time()))
        conn.commit()

    return templates.TemplateResponse("report.html", {
        "request": request,
        "checkout_id": checkout_id,
        "filename": file.filename,
        "s": summary,
        "flagged": flagged[:8],
        "flagged_total": len(flagged),
        "fraud": fraud[:5],
        "price": PRICE_DISPLAY,
        "pay_email": PAY_EMAIL,
        "dev_mode": DEV_MODE,
    })

@app.post("/pay", response_class=HTMLResponse)
async def pay(request: Request, checkout_id: str = Form(...), email: str = Form(...), phone: str = Form(...)):
    with closing(db()) as conn:
        order = conn.execute("SELECT * FROM orders WHERE checkout_id = ?", (checkout_id,)).fetchone()
    if not order:
        return templates.TemplateResponse("error.html", {
            "request": request, "msg": "Order not found.", "dev_mode": DEV_MODE,
        }, status_code=404)

    digits = re.sub(r"\D", "", phone)
    if digits.startswith("0"): digits = "254" + digits[1:]
    elif digits.startswith(("7","1")): digits = "254" + digits
    phone_intl = "+" + digits
    ref = "VC-" + checkout_id

    with closing(db()) as conn:
        conn.execute("UPDATE orders SET email=?, phone=?, reference=? WHERE checkout_id=?",
                     (email, phone_intl, ref, checkout_id))
        conn.commit()

    if DEV_MODE:
        with closing(db()) as conn:
            conn.execute("UPDATE orders SET paid = 1 WHERE checkout_id = ?", (checkout_id,))
            conn.commit()
        await send_csv_email(email, order["filename"], order["cleaned_csv"], {
            "total": order["total"], "flagged": order["flagged"], "fraud": order["fraud"],
        })
        return templates.TemplateResponse("waiting.html", {
            "request": request, "checkout_id": checkout_id, "dev_mode": DEV_MODE,
        })

    async with httpx.AsyncClient(timeout=30) as client:
        r = await client.post("https://api.paystack.co/charge",
            headers={"Authorization": "Bearer " + PAYSTACK_SECRET, "Content-Type": "application/json"},
            json={"email": email, "amount": PRICE_KOBO, "currency": "KES", "reference": ref,
                  "mobile_money": {"phone": phone_intl, "provider": "mpesa"}})
        data = r.json()

    if not data.get("status"):
        return templates.TemplateResponse("error.html", {
            "request": request, "msg": data.get("message", "Payment failed."), "dev_mode": DEV_MODE,
        }, status_code=400)

    return templates.TemplateResponse("waiting.html", {
        "request": request, "checkout_id": checkout_id, "dev_mode": DEV_MODE,
    })

@app.get("/order/{checkout_id}/status")
def order_status(checkout_id: str):
    with closing(db()) as conn:
        row = conn.execute("SELECT paid FROM orders WHERE checkout_id = ?", (checkout_id,)).fetchone()
    if not row: raise HTTPException(404, "not found")
    return {"paid": bool(row["paid"])}

@app.get("/download/{checkout_id}")
def download(checkout_id: str):
    with closing(db()) as conn:
        row = conn.execute("SELECT cleaned_csv, paid, filename FROM orders WHERE checkout_id = ?", (checkout_id,)).fetchone()
    if not row: raise HTTPException(404, "not found")
    if not row["paid"]: raise HTTPException(402, "payment not confirmed")
    return Response(content=row["cleaned_csv"], media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="cleaned_' + row["filename"] + '"'})

@app.post("/webhook/paystack")
async def paystack_webhook(request: Request):
    raw = await request.body()
    sig = request.headers.get("x-paystack-signature", "")
    expected = hmac.new(PAYSTACK_SECRET.encode(), raw, hashlib.sha512).hexdigest()
    if not hmac.compare_digest(sig, expected): raise HTTPException(401, "bad signature")
    event = json.loads(raw)
    if event.get("event") == "charge.success":
        ref = event.get("data", {}).get("reference", "")
        if ref.startswith("VC-"):
            cid = ref[3:]
            with closing(db()) as conn:
                conn.execute("UPDATE orders SET paid = 1 WHERE checkout_id = ?", (cid,))
                conn.commit()
                row = conn.execute("SELECT * FROM orders WHERE checkout_id = ?", (cid,)).fetchone()
            if row and row["email"]:
                try:
                    await send_csv_email(row["email"], row["filename"], row["cleaned_csv"], {
                        "total": row["total"], "flagged": row["flagged"], "fraud": row["fraud"],
                    })
                except Exception as e:
                    print("webhook email failed:", repr(e))
    return {"status": "ok"}

@app.get("/healthz")
def health(): return {"ok": True, "dev_mode": DEV_MODE, "brevo": bool(BREVO_API_KEY)}
